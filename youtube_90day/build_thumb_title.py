"""サムネイル＆タイトル設計ワークブック生成スクリプト（3本の解説動画の要点を整理）."""
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Meiryo"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=16, bold=True, color="1F3864")
F_SUB = Font(name=FONT, size=10, italic=True, color="595959")
F_EX = Font(name=FONT, size=10, italic=True, color="808080")
F_SEC = Font(name=FONT, size=12, bold=True, color="1F3864")
FILL_HEAD = PatternFill("solid", fgColor="1F3864")
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
FILL_EX = PatternFill("solid", fgColor="F2F2F2")
FILL_SEC = PatternFill("solid", fgColor="D9E1F2")
FILL_OK = PatternFill("solid", fgColor="C6EFCE")
FILL_NG = PatternFill("solid", fgColor="FFC7CE")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEGEND = "凡例：黄色＝記入欄 ／ 灰色の斜体＝記入例（上書き・削除OK） ／ 白＝自動計算"
SRC = "出典：①サムネ動画 ②タイトル動画 ③ブレイク分析動画 ／ ※＝動画にない、日本語向けの目安（作成者による）"

CURIOSITY = '"瞬間型,ストーリー型,結果型,変身型,新奇性型"'
STOPPERS = '"顔,見慣れたもの,大きな数字,お金,危険・動き,感情,鮮やかな色,美しさ"'
SCORE = '"1,2,3,4,5"'

wb = Workbook()


def setup(ws, title, subtitle, widths):
    ws["A1"], ws["A2"], ws["A3"] = title, subtitle, LEGEND
    ws["A1"].font, ws["A2"].font, ws["A3"].font = F_TITLE, F_SUB, F_SUB
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False


def header(ws, row, labels, col=1):
    for c, label in enumerate(labels, col):
        cell = ws.cell(row=row, column=c, value=label)
        cell.font, cell.fill, cell.alignment, cell.border = F_HEAD, FILL_HEAD, CENTER, BORDER
    ws.row_dimensions[row].height = 32


def put(ws, r, c, v=None, font=F_BASE, fill=None, align=WRAP, fmt=None):
    cell = ws.cell(row=r, column=c)
    if v is not None:
        cell.value = v
    cell.font, cell.border, cell.alignment = font, BORDER, align
    if fill:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    return cell


def inputs(ws, rows, cols, examples=()):
    """rows×cols を記入欄にし、先頭行から examples（行ごとのリスト）を記入例として入れる."""
    for r in rows:
        for c in cols:
            put(ws, r, c, fill=FILL_INPUT)
    for r, ex in zip(rows, examples):
        for c, v in zip(cols, ex):
            if v is not None:
                put(ws, r, c, v, font=F_EX, fill=FILL_EX)


def dv(ws, formula, rng):
    d = DataValidation(type="list", formula1=formula, allow_blank=True)
    ws.add_data_validation(d)
    d.add(rng)


def section(ws, row, text, ncols):
    ws.cell(row=row, column=1, value=text).font = F_SEC
    for c in range(1, ncols + 1):
        ws.cell(row=row, column=c).fill = FILL_SEC


# =====================================================================
# 00 原則まとめ
# =====================================================================
ws = wb.active
ws.title = "00_原則まとめ"
setup(ws, "サムネイル＆タイトル設計：原則まとめ", SRC, [22, 70, 10])
header(ws, 5, ["テーマ", "要点", "出典"])
principles = [
    ("【土台】", None, None),
    ("チャンネル戦略", "次の動画は、前の動画の視聴者の8割が興味を持つか（80%ルール）。視聴者層とフォーマットがぶれると勢いが次の動画に引き継がれない", "③"),
    ("動画レシピ", "個性・制作/編集の質・中身の価値のうち2つ以上で差別化。自分にしか出せない要素（経験・スキル・人柄）で真似されにくくする", "③"),
    ("企画が先", "サムネが思いつかない企画は、そもそも企画が弱い可能性。良いサムネは良い企画から生まれる", "①"),
    ("時間配分", "サムネ・タイトルは結果の半分を左右する。公開直前ではなく、制作に入る前に作る。数日かけてもよい", "①②"),
    ("【クリックの心理】", None, None),
    ("好奇心のギャップ", "「知っていること」と「知りたいこと」の差。サムネの8〜9割は心理、デザインは1〜2割", "①②"),
    ("サムネの5つの型", "瞬間型（リアクション直前）／ストーリー型（緊張・問い）／結果型（欲しい成果）／変身型（前→後）／新奇性型（見たことがない・意外）", "①"),
    ("タイトルの2つの作り方", "結末を伏せて引っぱる（オープンループ）／比べて戦わせる（対立）。どちらも「書かないこと」で好奇心を生む", "②"),
    ("スクロールストッパー", "顔・見慣れたもの・大きな数字・お金・危険や動き・感情・鮮やかな色・美しさ。使うのは1〜2個。詰め込むと釣りっぽくなりブランドを損なう", "①"),
    ("小さいチャンネルの顔", "知名度がないうちは顔の効果が弱い。顔だけに頼らず他のストッパーも使う", "①"),
    ("【デザイン（3つのC）】", None, None),
    ("中身", "好奇心のギャップに必要なものだけ。主役は1つ、他は主役を支える脇役", "①"),
    ("構図", "左右対称（主役を中央）か三分割（主役を左右1/3）。線・矢印で視線を主役へ。大きさ・ぼかし・奥行きで主役を最上位に", "①"),
    ("サムネの文字", "タイトルの繰り返しは不可。規模・背景・補足など情報を足すときだけ。英語なら5語以内。主役より目立たせない", "①"),
    ("コントラスト", "明度（明暗）・彩度（鮮やか/灰色）・色相（補色）で主役を浮かせる。同ジャンルの流行と逆を行くと目立つ", "①"),
    ("【タイトル】", None, None),
    ("型を借りる", "ゼロから考えない。伸びた型の構造だけ借りる（例：【期間】で【成果】を作った方法／【痛い結果】を招く【〇〇】のミス）", "②"),
    ("強い言葉", "「良い・コツ」より、感情・緊急性が伝わる言葉（台無し・誰も知らない・圧倒的 など）", "②"),
    ("長さ", "英語は50字未満・6〜10語。日本語はスマホで切れる前の前半約30字に要点を入れる", "②※"),
    ("サムネと分担", "サムネで見れば分かることはタイトルに書かない。2つでチームにする", "②"),
    ("【研究と検証】", None, None),
    ("アウトライヤー", "チャンネル平均より桁違いに伸びた動画。効いた要素がアイデア・タイトル・サムネのどれかを特定し、借りるのは1つだけ。丸ごと真似は埋もれる", "③"),
    ("案の数", "タイトルは20〜50案、サムネは下書き10〜15案。完成させるのは3案（公開用＋差し替え用2つ）", "①②"),
    ("公開後", "数時間〜48時間で判断。全部平均以上→そのまま／再生・表示が多くCTRだけ低い→外に広がっている証拠なのでそのまま／全部平均以下→差し替え", "①②"),
    ("心構え", "完璧なサムネはない。ガイドラインは破ってよいが、まず知ること。伸びなかった動画も学びになる。続けた人が勝つ", "①③"),
]
r = 6
for theme, point, src in principles:
    if point is None:
        section(ws, r, theme, 3)
    else:
        put(ws, r, 1, theme, font=F_BOLD)
        put(ws, r, 2, point)
        put(ws, r, 3, src, align=CENTER)
    r += 1

# =====================================================================
# 01 アウトライヤー分析
# =====================================================================
ws = wb.create_sheet("01_アウトライヤー分析")
setup(ws, "分析すべきポイント：伸びた動画を分解する",
      "登録者数に対して再生数が多い動画を集め、何が効いたかを分解する。倍率3倍以上を目安に「候補」表示（※動画の分析ツールの指標3を参考にした簡易版）",
      [4, 30, 10, 10, 8, 10, 11, 12, 12, 22, 14, 14, 12, 12, 16, 16, 24])
cols = ["No", "動画タイトル／URL", "登録者数", "再生数", "倍率", "判定", "一番効いた要素", "サムネの型",
        "タイトルの作り方", "タイトルの型（構造）", "強い言葉", "ストッパー", "主役", "構図",
        "コントラスト", "サムネ文字の役割", "自分が借りる1要素"]
header(ws, 5, cols)
ex = [
    ["IELTS独学で7.0取った人の勉強法（仮）", 8000, 95000, None, None, "サムネ", "結果型", "オープンループ",
     "【期間】で【成果】を取った方法", "独学・たった", "大きな数字", "スコア表", "三分割",
     "明度（暗い背景に白文字）", "規模（3ヶ月）", "スコア表を主役にする構図"],
    ["社会人が英語100時間やった結果（仮）", 30000, 400000, None, None, "アイデア", "変身型", "オープンループ",
     "【数字】時間やった結果", "結果", "顔", "本人の表情", "左右対称", "彩度（背景灰色）", "なし",
     "ビフォーアフターの企画"],
]
inputs(ws, range(6, 26), [2, 3, 4] + list(range(7, 18)))
for r, row in zip((6, 7), ex):
    for c, v in zip([2, 3, 4] + list(range(7, 18)), [row[0], row[1], row[2]] + row[5:]):
        put(ws, r, c, v, font=F_EX, fill=FILL_EX)
for r in range(6, 26):
    put(ws, r, 1, r - 5, align=CENTER)
    put(ws, r, 5, f'=IF(OR(C{r}="",D{r}="",C{r}=0),"",D{r}/C{r})', align=CENTER, fmt="0.0")
    put(ws, r, 6, f'=IF(E{r}="","",IF(E{r}>=3,"候補","－"))', align=CENTER)
    for c in (3, 4):
        ws.cell(row=r, column=c).number_format = "#,##0"
dv(ws, '"アイデア,タイトル,サムネ"', "G6:G25")
dv(ws, CURIOSITY, "H6:H25")
dv(ws, '"オープンループ,対立,両方,なし"', "I6:I25")
dv(ws, STOPPERS, "L6:L25")
dv(ws, '"左右対称,三分割"', "N6:N25")
dv(ws, '"規模,背景の説明,補足,見出し,なし"', "P6:P25")
ws.conditional_formatting.add("F6:F25", CellIsRule(operator="equal", formula=['"候補"'], fill=FILL_OK))
ws["A27"] = "見方：①平均より伸びた動画を探す → ②アイデア・タイトル・サムネのどれが効いたか1つ特定 → ③その1要素だけ自分のレシピに合わせて借りる（丸ごと真似はしない）"
ws["A27"].font = F_SUB
ws.freeze_panes = "C6"

# =====================================================================
# 02 作る手順
# =====================================================================
ws = wb.create_sheet("02_作る手順")
setup(ws, "サムネイルとタイトルを作る手順", "上から順に進める。動画の制作（台本・撮影）に入る前に STEP 7 まで終わらせるのが理想", [8, 14, 44, 44, 8, 8])
header(ws, 5, ["STEP", "段階", "やること", "考える問い", "出典", "済✔"])
steps = [
    ("0", "土台", "チャンネル戦略と動画レシピを確認する", "前の動画の視聴者の8割がこの動画に興味を持つ？ いつものフォーマット・レシピに乗っている？", "③"),
    ("1", "企画", "企画の強さを確かめる", "好奇心のギャップが作れる企画か？ サムネが思いつかないなら企画を見直す", "①"),
    ("2", "研究", "アウトライヤーを探して 01 に記録する", "登録者に対して桁違いに伸びた動画は？ ジャンルの外にもないか？", "③"),
    ("3", "研究", "効いた要素を1つだけ特定し、借りるものを決める", "アイデア・タイトル・サムネのどれが効いた？ 自分のレシピに合わせるとどうなる？", "③"),
    ("4", "心理", "好奇心の型と、伏せる情報を決める", "5つの型のどれ？ 視聴者が「答えを知りたい」と思う問いは何？", "①②"),
    ("5", "タイトル", "タイトルを20案以上出して採点する（03）", "結末を伏せている？ 対立がある？ 実績ある型？ 強い言葉？ 前半30字※で伝わる？", "②"),
    ("6", "サムネ", "サムネの下書きを10〜15案出す（絵は下手でOK）", "型・ストッパー・構図・コントラストを変えた、まったく別の案になっている？", "①"),
    ("7", "組み合わせ", "タイトル×サムネを3セットに絞る", "同じことを繰り返していない？ 2つで1つの好奇心になっている？", "①②"),
    ("8", "デザイン", "3セットのサムネを仕上げる（3つのC）", "主役は1つで一番目立つ？ 文字は情報を足している？ 10字以内※？", "①"),
    ("9", "テスト", "3つのテストをする（04）", "小さい表示で読める？ 競合と並べて目立つ？ 他人に2秒見せて伝わる？", "①"),
    ("10", "公開", "セットAで公開し、B・Cを待機させる", "差し替え用の2案は色違いではなく別コンセプトになっている？", "①"),
    ("11", "検証", "公開後数時間〜48時間で判定する（05）", "再生・表示・CTRはチャンネル平均と比べてどう？", "①②"),
    ("12", "学び", "結果を 05 と 01 に残し、次の動画に使う", "何が効いた？ 次に借りる型は？", "③"),
]
for i, s in enumerate(steps):
    r = 6 + i
    for c, v in enumerate(s, 1):
        put(ws, r, c, v, font=F_BOLD if c <= 2 else F_BASE, align=CENTER if c in (1, 2, 5) else WRAP)
    put(ws, r, 6, fill=FILL_INPUT, align=CENTER)
    ws.row_dimensions[r].height = 42
dv(ws, '"✔"', "F6:F18")
ws["A20"] = '=COUNTIF(F6:F18,"✔")&" / 13 STEP 完了"'
ws["A20"].font = F_BOLD

# =====================================================================
# 03 案出しシート
# =====================================================================
ws = wb.create_sheet("03_案出しシート")
setup(ws, "タイトル・サムネ案出しシート（1動画ぶん）", "このシートをコピーして動画ごとに使う。採点は1〜5。合計と順位は自動",
      [4, 40, 7, 30, 9, 9, 9, 9, 8, 6])
rows = [("動画のテーマ", "IELTSライティングを独学で6.5→7.0"),
        ("好奇心の型", "結果型"),
        ("伏せる情報（答え）", "添削なしで上げた具体的な方法"),
        ("借りる要素（01のNo）", "No1：スコア表を主役にするサムネ構図"),
        ("ストッパー（1〜2個）", "大きな数字（7.0）")]
for i, (k, v) in enumerate(rows):
    r = 4 + i
    put(ws, r, 1, k, font=F_BOLD)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    put(ws, r, 3, v, font=F_EX, fill=FILL_EX)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
dv(ws, CURIOSITY, "C5")

section(ws, 10, "タイトル案（20案以上。全部書き出してから採点）", 10)
header(ws, 11, ["No", "タイトル案", "字数", "スマホで見える部分（前半30字※）", "好奇心", "強い言葉", "分かりやすさ", "サムネと重複しない", "合計/20", "順位"])
titles = [
    ["IELTSライティング 独学で6.5→7.0にした添削ルーティン", None, None, 4, 3, 5, 4],
    ["添削なしでIELTSライティング7.0を取った、たった1つの習慣", None, None, 5, 4, 4, 5],
    ["IELTSライティングが6.5で止まる人がやっている3つのミス", None, None, 4, 4, 5, 4],
]
inputs(ws, range(12, 32), [2, 5, 6, 7, 8])
for r, t in zip(range(12, 15), titles):
    for c, v in zip([2, 5, 6, 7, 8], [t[0]] + t[3:]):
        put(ws, r, c, v, font=F_EX, fill=FILL_EX)
for r in range(12, 32):
    put(ws, r, 1, r - 11, align=CENTER)
    put(ws, r, 3, f'=IF(B{r}="","",LEN(B{r}))', align=CENTER)
    put(ws, r, 4, f'=IF(B{r}="","",LEFT(B{r},30))')
    put(ws, r, 9, f'=IF(COUNT(E{r}:H{r})=4,SUM(E{r}:H{r}),"")', font=F_BOLD, align=CENTER)
    put(ws, r, 10, f'=IF(I{r}="","",RANK(I{r},$I$12:$I$31))', align=CENTER)
dv(ws, SCORE, "E12:H31")
ws.conditional_formatting.add("C12:C31", CellIsRule(operator=">", formula=["50"], fill=FILL_NG))

section(ws, 33, "サムネ下書き（10〜15案。色違いではなく別コンセプトで）", 10)
header(ws, 34, ["No", "コンセプト（何が映っている？）", "型", "主役／ストッパー", "文字", "字数", "好奇心", "2秒で伝わる", "目立つ", "合計/15"])
thumbs = [
    ["左に赤ペンだらけの6.5答案、右に7.0のスコア表", "変身型", "スコア表／大きな数字", "独学"],
    ["7.0のスコア表を持って驚く顔、背景は暗く", "結果型", "スコア表／感情", "添削ゼロ"],
]
inputs(ws, range(35, 50), [2, 3, 4, 5, 7, 8, 9])
for r, t in zip((35, 36), thumbs):
    for c, v in zip([2, 3, 4, 5, 7, 8, 9], t + [5, 4, 4] if r == 35 else t + [4, 5, 3]):
        put(ws, r, c, v, font=F_EX, fill=FILL_EX)
for r in range(35, 50):
    put(ws, r, 1, r - 34, align=CENTER)
    put(ws, r, 6, f'=IF(E{r}="","",LEN(E{r}))', align=CENTER)
    put(ws, r, 10, f'=IF(COUNT(G{r}:I{r})=3,SUM(G{r}:I{r}),"")', font=F_BOLD, align=CENTER)
dv(ws, CURIOSITY, "C35:C49")
dv(ws, SCORE, "G35:I49")
ws.conditional_formatting.add("F35:F49", CellIsRule(operator=">", formula=["10"], fill=FILL_NG))

section(ws, 51, "3セットに絞る（A＝公開用、B・C＝差し替え用）", 10)
header(ws, 52, ["組", "タイトル", "", "サムネ", "", "", "役割分担（タイトルが言うこと／サムネが見せること）", "", "", ""])
for i, k in enumerate("ABC"):
    r = 53 + i
    put(ws, r, 1, k, font=F_BOLD, align=CENTER)
    for c1, c2 in ((2, 3), (4, 6), (7, 10)):
        put(ws, r, c1, fill=FILL_INPUT)
        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
put(ws, 53, 2, "添削なしでIELTSライティング7.0を取った、たった1つの習慣", font=F_EX, fill=FILL_EX)
put(ws, 53, 4, "No1：6.5答案→7.0スコア表", font=F_EX, fill=FILL_EX)
put(ws, 53, 7, "タイトル＝方法を伏せる／サムネ＝変化の大きさを見せる", font=F_EX, fill=FILL_EX)
ws.freeze_panes = "A12"

# =====================================================================
# 04 公開前チェック
# =====================================================================
ws = wb.create_sheet("04_公開前チェック")
setup(ws, "公開前チェックリスト", "全部「はい」を目指す。例外はOK（ガイドラインは破ってよいが、理由を持つ）", [14, 60, 10, 8])
header(ws, 5, ["分類", "チェック項目", "はい/いいえ", "出典"])
checks = [
    ("土台", "前の動画の視聴者の8割が興味を持つ内容か", "③"),
    ("土台", "いつものフォーマット・動画レシピに乗っているか", "③"),
    ("企画", "動画の制作前にサムネとタイトルを考えたか", "①"),
    ("企画", "アウトライヤーから借りた要素は1つだけか（丸ごと真似していない）", "③"),
    ("心理", "視聴者が答えを知りたくなる問いが1つある", "①②"),
    ("心理", "好奇心の型（5つのどれか）が決まっている", "①"),
    ("心理", "ストッパーは1〜2個で、詰め込みすぎていない", "①"),
    ("心理", "自分のブランドらしさを損なっていない（釣りっぽくない）", "①"),
    ("タイトル", "20案以上出して選んだ", "②"),
    ("タイトル", "結末を伏せている、または対立がある", "②"),
    ("タイトル", "実績のある型の構造を使っている", "②"),
    ("タイトル", "弱い言葉ではなく、感情が動く言葉がある", "②"),
    ("タイトル", "前半約30字で要点が伝わる※", "②※"),
    ("サムネ", "下書きを10案以上出して選んだ", "①"),
    ("サムネ", "主役は1つで、一番目立っている", "①"),
    ("サムネ", "好奇心に不要な要素を入れていない", "①"),
    ("サムネ", "構図（左右対称か三分割）が決まっていて、視線が主役に向かう", "①"),
    ("サムネ", "明度・彩度・色相のどれかで主役が浮いている", "①"),
    ("サムネ", "文字はタイトルの繰り返しではなく、情報を足している", "①"),
    ("サムネ", "文字は10字以内※で、主役より目立っていない", "①※"),
    ("組み合わせ", "タイトルとサムネで同じことを言っていない", "①②"),
    ("テスト", "明瞭さ：関連動画欄の小さい表示でも全部読める", "①"),
    ("テスト", "コントラスト：競合の動画と並べても目立つ", "①"),
    ("テスト", "チラ見：他人に2秒見せて、何の動画か伝わった", "①"),
    ("準備", "差し替え用に別コンセプトの2セットを用意した", "①②"),
    ("準備", "公開後に確認する時間（数時間後・24〜48時間後）を決めた", "①②"),
]
for i, (k, item, src) in enumerate(checks):
    r = 6 + i
    put(ws, r, 1, k, font=F_BOLD, align=CENTER)
    put(ws, r, 2, item)
    put(ws, r, 3, fill=FILL_INPUT, align=CENTER)
    put(ws, r, 4, src, align=CENTER)
last = 5 + len(checks)
dv(ws, '"はい,いいえ"', f"C6:C{last}")
ws.conditional_formatting.add(f"C6:C{last}", CellIsRule(operator="equal", formula=['"いいえ"'], fill=FILL_NG))
ws.conditional_formatting.add(f"C6:C{last}", CellIsRule(operator="equal", formula=['"はい"'], fill=FILL_OK))
ws["B4"] = f'="達成率："&TEXT(COUNTIF(C6:C{last},"はい")/{len(checks)},"0%")&"（はい "&COUNTIF(C6:C{last},"はい")&" / {len(checks)}）"'
ws["B4"].font = F_BOLD
for r in (6, 7, 8, 9, 10):
    put(ws, r, 3, "はい", font=F_EX, fill=FILL_EX, align=CENTER)
ws.freeze_panes = "A6"

# =====================================================================
# 05 公開後の検証
# =====================================================================
ws = wb.create_sheet("05_公開後の検証")
setup(ws, "公開後の検証と差し替え判定",
      "同じタイミング（例：公開24時間後）の数字をチャンネル平均と比べる。判定は自動。CTRは小数で（5.2%→0.052）",
      [4, 28, 6, 10, 10, 8, 10, 10, 8, 20, 24])
header(ws, 5, ["No", "動画タイトル", "セット", "再生数", "表示回数", "CTR", "平均 再生", "平均 表示", "平均CTR", "判定", "差し替え後の結果・学び"])
posts = [
    ["添削なしでIELTSライティング7.0…", "A", 900, 15000, 0.045, 1200, 20000, 0.05, None, "Bに差し替え→CTR 6.1%"],
    ["平日30分しかない社会人のIELTS…", "A", 3000, 90000, 0.035, 1200, 20000, 0.05, None, "外に広がり中、そのまま"],
]
inputs(ws, range(6, 21), [2, 3, 4, 5, 6, 7, 8, 9, 11])
for r, p in zip((6, 7), posts):
    for c, v in zip([2, 3, 4, 5, 6, 7, 8, 9, 11], p[:8] + [p[9]]):
        put(ws, r, c, v, font=F_EX, fill=FILL_EX)
for r in range(6, 21):
    put(ws, r, 1, r - 5, align=CENTER)
    for c in (4, 5, 7, 8):
        ws.cell(row=r, column=c).number_format = "#,##0"
    for c in (6, 9):
        ws.cell(row=r, column=c).number_format = "0.0%"
    put(ws, r, 10,
        f'=IF(COUNT(D{r}:I{r})<6,"",IF(AND(D{r}>=G{r},E{r}>=H{r},F{r}>=I{r}),"成功：そのまま",'
        f'IF(AND(D{r}>=G{r},E{r}>=H{r},F{r}<I{r}),"外に拡散中：そのまま",'
        f'IF(AND(D{r}<G{r},E{r}<H{r},F{r}<I{r}),"差し替え検討","様子見"))))',
        font=F_BOLD, align=CENTER)
dv(ws, '"A,B,C"', "C6:C20")
ws.conditional_formatting.add("J6:J20", CellIsRule(operator="equal", formula=['"差し替え検討"'], fill=FILL_NG))
ws["A23"] = "判定ルール（①②）：全部平均以上→そのまま／再生・表示は平均以上でCTRだけ低い→普段の視聴者の外に広がっているのでそのまま／全部平均以下→別セットに差し替え／それ以外→様子見"
ws["A23"].font = F_SUB
ws.freeze_panes = "C6"

out = "youtube_90day/YouTubeサムネ・タイトル設計シート.xlsx"
wb.save(out)
print("saved", out)
