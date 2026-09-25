"""90日間YouTubeチャレンジ ワークブック生成スクリプト."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter

FONT = "Meiryo"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=16, bold=True, color="1F3864")
F_SUB = Font(name=FONT, size=10, italic=True, color="595959")
F_EX = Font(name=FONT, size=10, italic=True, color="808080")
F_SEC = Font(name=FONT, size=12, bold=True, color="1F3864")

FILL_HEAD = PatternFill("solid", fgColor="1F3864")
FILL_P1 = PatternFill("solid", fgColor="DDEBF7")
FILL_P2 = PatternFill("solid", fgColor="FCE4D6")
FILL_P3 = PatternFill("solid", fgColor="E2EFDA")
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
FILL_EX = PatternFill("solid", fgColor="F2F2F2")
FILL_SEC = PatternFill("solid", fgColor="D9E1F2")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

LEGEND = "凡例：黄色セル＝あなたが記入する欄 ／ 灰色の斜体行＝記入例（上書き・削除OK） ／ 白セル＝自動計算"

wb = Workbook()


def setup(ws, title, subtitle, widths):
    ws["A1"] = title
    ws["A1"].font = F_TITLE
    ws["A2"] = subtitle
    ws["A2"].font = F_SUB
    ws["A3"] = LEGEND
    ws["A3"].font = F_SUB
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False


def header(ws, row, labels, fill=FILL_HEAD):
    for c, label in enumerate(labels, 1):
        cell = ws.cell(row=row, column=c, value=label)
        cell.font = F_HEAD
        cell.fill = fill
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[row].height = 32


def style(cell, font=F_BASE, fill=None, align=WRAP):
    cell.font = font
    cell.border = BORDER
    cell.alignment = align
    if fill:
        cell.fill = fill


def section(ws, row, text, ncols):
    ws.cell(row=row, column=1, value=text).font = F_SEC
    for c in range(1, ncols + 1):
        ws.cell(row=row, column=c).fill = FILL_SEC
    ws.row_dimensions[row].height = 22


def input_block(ws, first_row, n_rows, cols, example=None):
    """example: list of values for the first row (rendered as example)."""
    for r in range(first_row, first_row + n_rows):
        for c in cols:
            style(ws.cell(row=r, column=c), fill=FILL_INPUT)
    if example:
        for c, v in zip(cols, example):
            cell = ws.cell(row=first_row, column=c, value=v)
            style(cell, font=F_EX, fill=FILL_EX)


def add_list_validation(ws, formula, rng):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(rng)


# =====================================================================
# 0. ロードマップ
# =====================================================================
ws = wb.active
ws.title = "00_ロードマップ"
setup(ws, "90日間YouTubeチャレンジ ロードマップ",
      "12週間・3フェーズ。基準ペースは週1本（普段のペースに合わせて調整）。各週の「使うシート」を開いて進める。",
      [8, 16, 22, 48, 40, 22, 14, 14])
header(ws, 5, ["週", "フェーズ", "テーマ", "やること（ゴール）", "完了の目安", "使うシート", "予定開始日", "完了✔"])
roadmap = [
    ("W1", "1 土台作り", "決意表明", "目標を書き出し、公開宣言する（コミュニティ投稿・ショート・X・Instagram・Discord等。視聴者がいなければ紙＋友人/家族）", "宣言を投稿した", "01_決意表明"),
    ("W2", "1 土台作り", "審美眼を広げる", "クリエイター目線で30本以上研究。ニッチ外・多フォーマットも見る。スワイプファイル作成開始", "30本分析・スワイプファイルができた", "02_スワイプファイル"),
    ("W3", "1 土台作り", "動画レシピ開発", "5要素（ターゲット/価値/フォーマット/雰囲気/差別化）でテスト用レシピを1〜3個書く", "レシピ1〜3個", "03_動画レシピ"),
    ("W4", "1 土台作り", "アイデア創出", "3基準（検証可能性/再生ポテンシャル/実現可能性）を満たすアイデアを8本出す", "8本決定", "04_アイデア8本"),
    ("W5", "2 スプリント", "2回目の公開宣言＋制作開始", "スプリント宣言。投稿頻度を2倍に。各動画で4ステップ（学習目標→制作→振り返り→共有）", "動画1〜2本公開", "05_スプリント管理 / 06_制作シート"),
    ("W6", "2 スプリント", "スプリント継続", "とにかく完成させて公開。投稿ごとに20〜30分の振り返り", "累計4本", "05〜07"),
    ("W7", "2 スプリント", "スプリント継続", "3本目あたりで失速しやすい。想定内と割り切る", "累計6本", "05〜07"),
    ("W8", "2 スプリント", "スプリント継続", "振り返りのパターンを見る", "累計8本", "05〜07"),
    ("W9", "2 スプリント", "予備週", "遅れの取り戻し用バッファ。最後の振り返り・コミュニティ更新", "8本完了（本数より学び重視）", "05〜07"),
    ("W10", "3 渾身の1本", "特大ヒット作の制作", "スプリントの「もっと時間があったら」を全投入して渾身の1本を作る", "企画・台本・撮影", "08_渾身の1本"),
    ("W11", "3 渾身の1本", "特大ヒット作の公開", "品質を極限まで追求し公開（※W10-11配分は目安）", "公開", "08_渾身の1本"),
    ("W12", "3 渾身の1本", "最終振り返り", "全振り返りを見直し、量vs質の感覚を比較して今後の投稿スケジュールと黄金律を決める", "今後の方針決定", "09_最終振り返り"),
]
fills = {"1": FILL_P1, "2": FILL_P2, "3": FILL_P3}
for i, row in enumerate(roadmap):
    r = 6 + i
    for c, v in enumerate(row, 1):
        style(ws.cell(row=r, column=c, value=v), fill=fills[row[1][0]])
    ws.cell(row=r, column=1).alignment = CENTER
    # 予定開始日: 開始日から自動
    d = ws.cell(row=r, column=7, value=f"=IF($C$20=\"\",\"\",$C$20+{7 * i})")
    style(d, align=CENTER)
    d.number_format = "yyyy/mm/dd"
    style(ws.cell(row=r, column=8), fill=FILL_INPUT, align=CENTER)
    ws.row_dimensions[r].height = 45
add_list_validation(ws, '"✔"', "H6:H17")

ws["A19"] = "設定"
ws["A19"].font = F_SEC
ws["A20"] = "チャレンジ開始日"
ws["A20"].font = F_BOLD
style(ws["C20"], fill=FILL_INPUT, align=CENTER)
ws["C20"].number_format = "yyyy/mm/dd"
ws["D20"] = "← 日付を入れると「予定開始日」が自動で入ります"
ws["D20"].font = F_SUB
ws["A21"] = "普段の投稿ペース（本/週）"
ws["A21"].font = F_BOLD
style(ws["C21"], fill=FILL_INPUT, align=CENTER)
ws["C21"] = 1
ws["C21"].font = Font(name=FONT, size=10, color="0000FF")
ws["D21"] = "← 動画の基準は週1本。隔週なら 0.5"
ws["D21"].font = F_SUB
ws["A22"] = "スプリント中の目標ペース（本/週）"
ws["A22"].font = F_BOLD
style(ws["C22"], align=CENTER)
ws["C22"] = "=C21*2"
ws["D22"] = "← 普段の2倍（自動）"
ws["D22"].font = F_SUB
ws["A23"] = "8本完了までの目安（週）"
ws["A23"].font = F_BOLD
style(ws["C23"], align=CENTER)
ws["C23"] = "=IF(C22>0,ROUNDUP(8/C22,0),\"\")"

ws["A25"] = "進捗ダッシュボード（自動）"
ws["A25"].font = F_SEC
dash = [
    ("完了した週", "=COUNTIF(H6:H17,\"✔\")&\" / 12\""),
    ("スワイプファイル件数（目標30）", "=COUNTA('02_スワイプファイル'!B7:B46)"),
    ("動画レシピ数（目標1〜3）", "=COUNTA('03_動画レシピ'!C6:E6)"),
    ("アイデア数（目標8）", "=COUNTA('04_アイデア8本'!B7:B16)"),
    ("スプリント公開本数（目標8）", "='05_スプリント管理'!C4"),
    ("振り返り記入済み本数", "=COUNTA('07_投稿後振り返り'!B7:B16)"),
]
for i, (label, f) in enumerate(dash):
    r = 26 + i
    ws.cell(row=r, column=1, value=label).font = F_BOLD
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    style(ws.cell(row=r, column=3, value=f), font=F_BOLD, align=CENTER)
ws.freeze_panes = "A6"

# =====================================================================
# 1. 決意表明
# =====================================================================
ws = wb.create_sheet("01_決意表明")
setup(ws, "W1 決意表明", "目標を書くと達成率が上がり、人に共有するとさらに上がる（ゲイル・マシューズ博士の研究）。挫折の原因は下手さではなく「諦めること」。",
      [36, 70])
qs = [
    ("90日後の具体的な目標（数値で）", "例：登録者1,000人・総再生4,000時間で収益化"),
    ("なぜ今それをやるのか（本気度）", "例：今年中に仕事と両立できる柱にしたい"),
    ("止まりそうな時に読み返す一言", "例：完璧より公開。出さない動画からは何も学べない"),
    ("予想される挫折ポイントと対策", "例：残業週→撮影は土曜午前に固定"),
    ("宣言する場所（1回目）", "例：コミュニティ投稿＋X"),
    ("宣言文（下書き）", "例：今日から90日間チャレンジを始めます。12週間で…"),
    ("宣言した日", ""),
    ("宣言を見せた相手／反応", ""),
    ("【W5用】2回目の宣言場所", "例：ショート動画で「4週間スプリントで投稿2倍にします」"),
    ("【W5用】2回目の宣言文", ""),
    ("【W5用】2回目を宣言した日", ""),
]
header(ws, 5, ["考えること", "あなたの答え"])
for i, (q, ex) in enumerate(qs):
    r = 6 + i
    style(ws.cell(row=r, column=1, value=q), font=F_BOLD)
    c = ws.cell(row=r, column=2, value=ex if ex else None)
    style(c, font=F_EX if ex else F_BASE, fill=FILL_INPUT)
    if "日" in q and q.endswith("日"):
        c.number_format = "yyyy/mm/dd"
    ws.row_dimensions[r].height = 36

# =====================================================================
# 2. スワイプファイル
# =====================================================================
ws = wb.create_sheet("02_スワイプファイル")
setup(ws, "W2 審美眼を広げる：動画研究＆スワイプファイル",
      "消費者ではなくクリエイターとして見る。「面白かったか」ではなく「なぜ」を書く。目標30本以上、ニッチ外・多フォーマットも。",
      [5, 30, 18, 14, 14, 30, 30, 30, 30, 26, 26, 30, 10])
ws["F4"] = "研究本数："
ws["F4"].font = F_BOLD
ws["G4"] = "=COUNTA(B7:B46)&\" / 30本\""
ws["G4"].font = F_BOLD
cols = ["No", "動画タイトル / URL", "チャンネル", "フォーマット", "ニッチ内/外",
        "なぜクリックした？（タイトル・サムネ）", "30秒後も見続けた理由（フック）",
        "音楽・照明・編集で湧いた感情", "退屈・離脱したタイミングときっかけ",
        "盗みたい要素（スワイプ）", "自分の動画への応用案", "雰囲気（一言）", "重要度\n(1-5)"]
header(ws, 6, cols)
FORMATS = '"Vlog,チャレンジ,チュートリアル,解説,分析,リスト/ランキング,ドキュメンタリー,インタビュー,リアクション,ショート,その他"'
example = ["https://youtube.com/... 「3ヶ月で英語が話せた方法」", "〇〇English", "解説", "ニッチ内",
           "数字＋ビフォーアフター。サムネの表情が強い", "冒頭で結果を見せてから過程を約束",
           "静かなLo-fiで落ち着く、自然光で親近感", "中盤の機材説明（3:40）で離脱しそうに",
           "結果→過程の順で見せる冒頭構成", "IELTSスコア公開から始める", "親しみ・前向き", 5]
input_block(ws, 7, 40, range(2, 14), example)
for r in range(7, 47):
    c = ws.cell(row=r, column=1, value=r - 6)
    style(c, align=CENTER)
add_list_validation(ws, FORMATS, "D7:D46")
add_list_validation(ws, '"ニッチ内,ニッチ外"', "E7:E46")
add_list_validation(ws, '"1,2,3,4,5"', "M7:M46")
ws.freeze_panes = "C7"

# =====================================================================
# 3. 動画レシピ
# =====================================================================
ws = wb.create_sheet("03_動画レシピ")
setup(ws, "W3 独自の動画レシピ（テスト用に1〜3個）",
      "最終決定ではなく「実験するための仮説」。強み（得意なこと）とスワイプファイルで繰り返し出てくる要素から組み立てる。",
      [18, 46, 34, 34, 34])
header(ws, 5, ["要素", "自分への問い", "レシピA", "レシピB", "レシピC"])
style(ws.cell(row=6, column=1, value="レシピ名"), font=F_BOLD)
style(ws.cell(row=6, column=2, value="短い呼び名（アイデア・振り返りシートで使う）"))
recipe = [
    ("① ターゲット層", "誰のため？何に興味があり、何に悩んでいる？普段どんな動画を見ている？"),
    ("② 価値", "見た人は何を得る？行動の後押し／理解が深まる／孤独が和らぐ／娯楽／インスピレーション"),
    ("③ フォーマット", "その価値をどう届ける？（Vlog・解説・チュートリアル・チャレンジ等から強みに合う1〜2個）"),
    ("④ 雰囲気（バイブス）", "見終わった後どんな気分にさせたい？（居心地がいい・スリル・刺激的・効率的…）"),
    ("⑤ 差別化要因", "なぜ他の人ではなく自分の動画を見るべき？（独自スキル／アクセス権／深い洞察／個性）"),
]
for i, (el, q) in enumerate(recipe):
    r = 7 + i
    style(ws.cell(row=r, column=1, value=el), font=F_BOLD)
    style(ws.cell(row=r, column=2, value=q))
    ws.row_dimensions[r].height = 60
input_block(ws, 6, 6, [3, 4, 5])
ex = ["IELTS独学ラボ", "IELTS 6.5を目指す社会人（独学・時間がない）",
      "スコアを上げる具体的な手順がわかる", "解説＋画面録画チュートリアル",
      "効率的・安心感", "自分が実際に独学で上げた記録と失敗例を全公開"]
for i, v in enumerate(ex):
    style(ws.cell(row=6 + i, column=3, value=v), font=F_EX, fill=FILL_EX)
ws.row_dimensions[6].height = 24

section(ws, 13, "レシピを作る前に考えること", 5)
pre = [
    ("自分が生まれつき得意なことは？", ""),
    ("スワイプファイルで何度も出てくるフォーマットは？", ""),
    ("スワイプファイルで何度も出てくる雰囲気は？", ""),
    ("他の人にない経験・アクセス・スキルは？", ""),
]
for i, (q, _) in enumerate(pre):
    r = 14 + i
    style(ws.cell(row=r, column=1, value=q), font=F_BOLD)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
    style(ws.cell(row=r, column=3), fill=FILL_INPUT)
    ws.row_dimensions[r].height = 30

# =====================================================================
# 4. アイデア8本
# =====================================================================
ws = wb.create_sheet("04_アイデア8本")
setup(ws, "W4 アイデア創出：3基準でスコアリングして8本選ぶ",
      "各基準を1〜5で評価。合計と順位は自動。複数レシピがある場合は各レシピ最低1本。大規模・複雑なネタは「渾身の1本」候補へ回す。",
      [5, 36, 16, 34, 30, 12, 12, 12, 10, 8, 30, 14])
cols = ["No", "動画アイデア（仮タイトル）", "使うレシピ", "検証したいこと（実験の問い）",
        "参考にした実績コンテンツ（URL・再生数）", "検証可能性\n(1-5)", "再生\nポテンシャル\n(1-5)",
        "実現可能性\n(2倍速で作れる？)(1-5)", "合計\n(/15)", "順位", "簡略化するなら？", "判定"]
header(ws, 6, cols)
example = ["IELTSライティング 独学で6.5→7.0にした添削ルーティン", "IELTS独学ラボ",
           "画面録画中心のフォーマットで維持率は保てるか？", "類似動画 45万回（登録者3万ch）", 4, 5, 4]
input_block(ws, 7, 10, range(2, 9), example)
input_block(ws, 7, 10, [11])
for r in range(7, 17):
    style(ws.cell(row=r, column=1, value=r - 6), align=CENTER)
    s = ws.cell(row=r, column=9, value=f'=IF(COUNT(F{r}:H{r})=3,SUM(F{r}:H{r}),"")')
    style(s, font=F_BOLD, align=CENTER)
    k = ws.cell(row=r, column=10, value=f'=IF(I{r}="","",RANK(I{r},$I$7:$I$16))')
    style(k, align=CENTER)
    j = ws.cell(row=r, column=12,
                value=f'=IF(I{r}="","",IF(H{r}<=2,"簡略化 or 渾身候補へ",IF(I{r}>=11,"採用候補","要ブラッシュアップ")))')
    style(j, align=CENTER)
    ws.row_dimensions[r].height = 36
add_list_validation(ws, '"1,2,3,4,5"', "F7:H16")
add_list_validation(ws, "'03_動画レシピ'!$C$6:$E$6", "C7:C16")
ws.conditional_formatting.add("I7:I16", CellIsRule(operator=">=", formula=["11"], fill=PatternFill("solid", fgColor="C6EFCE")))
ws.conditional_formatting.add("H7:H16", CellIsRule(operator="<=", formula=["2"], fill=PatternFill("solid", fgColor="FFC7CE")))
ws["A18"] = "基準の問い"
ws["A18"].font = F_SEC
notes = [
    "検証可能性：このアイデアでレシピの何がわかる？（このフォーマットを試したら？この雰囲気を強調したら？このターゲットにこの価値を出したら？）",
    "再生ポテンシャル：目標再生数を既に出しているタイトル/サムネ/コンセプト（自他・ニッチ内外）を自分のレシピに通したか？",
    "実現可能性：普段の2倍のペースで台本→撮影→編集→公開できる？無理なら簡略化できる？",
    "※8本を超えて出してOK（10行用意）。上位8本を 05_スプリント管理 に転記。",
]
for i, n in enumerate(notes):
    ws.cell(row=19 + i, column=1, value=n).font = F_BASE
ws.freeze_panes = "C7"

# =====================================================================
# 5. スプリント管理
# =====================================================================
ws = wb.create_sheet("05_スプリント管理")
setup(ws, "W5〜W9 スプリントモード：8本の進捗管理",
      "投稿頻度2倍・8本・5週間（1週は予備）。本数より「1本ごとの学び」。遅れても前進を続ける。",
      [5, 36, 16, 36, 13, 13, 14, 12, 12, 12, 12, 30])
ws["A4"] = "公開本数"
ws["A4"].font = F_BOLD
ws["C4"] = '=COUNTIF(G7:G14,"公開済")'
ws["C4"].font = F_BOLD
ws["D4"] = '=C4&" / 8本（"&TEXT(C4/8,"0%")&"）"'
ws["D4"].font = F_BOLD
ws["E4"] = "振り返り済"
ws["E4"].font = F_BOLD
ws["F4"] = '=COUNTIF(J7:J14,"✔")&" / "&C4'
ws["F4"].font = F_BOLD
cols = ["#", "動画タイトル", "レシピ", "学習目標（この動画で学ぶ1つのこと）", "公開予定日", "実際の公開日",
        "ステータス", "予定との差(日)", "学習目標設定", "振り返り記入", "コミュニティ共有", "メモ・詰まったこと"]
header(ws, 6, cols)
example = ["IELTSライティング 独学で6.5→7.0にした添削ルーティン", "IELTS独学ラボ",
           "スワイプの「結果→過程」冒頭で30秒維持率が上がるか", None, None, "制作中", None, "✔", None, None, "台本2時間で書けた"]
input_block(ws, 7, 8, [2, 3, 4, 5, 6, 7, 9, 10, 11, 12], None)
for c, v in zip([2, 3, 4, 7, 9, 12], [example[0], example[1], example[2], example[5], example[7], example[10]]):
    style(ws.cell(row=7, column=c, value=v), font=F_EX, fill=FILL_EX)
for r in range(7, 15):
    style(ws.cell(row=r, column=1, value=r - 6), align=CENTER)
    for c in (5, 6):
        ws.cell(row=r, column=c).number_format = "yyyy/mm/dd"
    d = ws.cell(row=r, column=8, value=f'=IF(OR(E{r}="",F{r}=""),"",F{r}-E{r})')
    style(d, align=CENTER)
    ws.row_dimensions[r].height = 36
add_list_validation(ws, '"未着手,企画,台本,撮影,編集,公開済"', "G7:G14")
add_list_validation(ws, '"✔"', "I7:K14")
add_list_validation(ws, "'03_動画レシピ'!$C$6:$E$6", "C7:C14")
ws.conditional_formatting.add("G7:G14", CellIsRule(operator="equal", formula=['"公開済"'], fill=PatternFill("solid", fgColor="C6EFCE")))
ws.conditional_formatting.add("H7:H14", CellIsRule(operator=">", formula=["0"], font=Font(name=FONT, color="C00000", bold=True)))

section(ws, 16, "各動画の4ステップ（毎回これを回す）", 12)
steps = [
    ("1 学習目標を設定", "台本・撮影の前に「この動画で学ぶ1つのこと」を決める（新フォーマット／タイトル構成／サムネ／編集の簡略化など）。→ D列"),
    ("2 動画を制作", "とにかく完成させる。考えすぎない・完璧を求めない。未完成に感じても約束したペースで公開。→ 06_制作シート"),
    ("3 振り返る", "公開後20〜30分で振り返りシートを記入。絶対に飛ばさない（最重要）。→ 07_投稿後振り返り"),
    ("4 コミュニティ共有（任意・推奨）", "「8本中3本目完了！」や振り返りの気づきを共有。投稿ごと or 週1。"),
]
for i, (s, d) in enumerate(steps):
    r = 17 + i
    style(ws.cell(row=r, column=1, value=s), font=F_BOLD)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    style(ws.cell(row=r, column=3, value=d))
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=12)
    ws.row_dimensions[r].height = 30
ws.cell(row=22, column=1, value="心構え：3本目あたりで新鮮さが消え、再生が伸びない・遅れる・疑問が湧くのは普通。1回目9本、2回目は4本でも効果があった。").font = F_SUB
ws.freeze_panes = "C7"

# =====================================================================
# 6. 制作シート（1本ごと）
# =====================================================================
ws = wb.create_sheet("06_制作シート")
setup(ws, "制作シート：1本ごとの企画〜公開チェック",
      "スプリント8本＋渾身の1本（#9）。各動画の列に記入。短い制作期間でも抜けがないようにする。",
      [34] + [24] * 9)
header(ws, 5, ["項目"] + [f"動画 #{i}" for i in range(1, 9)] + ["渾身 #9"])
for c in range(2, 11):
    col = get_column_letter(c)
    ws.cell(row=5, column=c).value = f"動画 #{c - 1}" if c < 10 else "渾身の1本"
items = [
    ("【企画】", None),
    ("動画タイトル（←05から自動）", "title"),
    ("学習目標（←05から自動）", "goal"),
    ("ターゲットが抱える悩み", None),
    ("視聴者が得る価値（1文）", None),
    ("タイトル案A", None),
    ("タイトル案B", None),
    ("タイトル案C", None),
    ("採用タイトル", None),
    ("サムネイル案（構図・文字・表情）", None),
    ("冒頭30秒のフック", None),
    ("構成（起承転結／見出し）", None),
    ("参考スワイプNo.", None),
    ("【制作】", None),
    ("台本 完了", "chk"),
    ("撮影 完了", "chk"),
    ("編集 完了", "chk"),
    ("サムネ 完了", "chk"),
    ("概要欄・タグ・チャプター", "chk"),
    ("#ハッシュタグ・宣言/タグ付け", "chk"),
    ("公開", "chk"),
    ("制作にかかった時間(h)", "num"),
    ("チェック完了率", "pct"),
]
title_src = {c: f"='05_スプリント管理'!B{c + 5}" for c in range(2, 10)}
goal_src = {c: f"='05_スプリント管理'!D{c + 5}" for c in range(2, 10)}
chk_first = None
for i, (label, kind) in enumerate(items):
    r = 6 + i
    if kind is None and label.startswith("【"):
        section(ws, r, label, 10)
        continue
    style(ws.cell(row=r, column=1, value=label), font=F_BOLD)
    for c in range(2, 11):
        cell = ws.cell(row=r, column=c)
        if kind == "title" and c < 10:
            cell.value = f'=IF({title_src[c][1:]}="","",{title_src[c][1:]})'
            style(cell)
        elif kind == "goal" and c < 10:
            cell.value = f'=IF({goal_src[c][1:]}="","",{goal_src[c][1:]})'
            style(cell)
        elif kind == "pct":
            pass  # 後で数式を入れる
        else:
            style(cell, fill=FILL_INPUT, align=CENTER if kind in ("chk", "num") else WRAP)
    if kind == "chk":
        chk_first = chk_first or r
    if kind in ("title", "goal"):
        pass
    ws.row_dimensions[r].height = 30 if kind not in ("chk", "num", "pct") else 18
    if label == "【制作】":
        pass
# compute chk rows after loop is too late for pct; recompute now
chk_rows = [6 + i for i, (_, k) in enumerate(items) if k == "chk"]
pct_row = 6 + [i for i, (_, k) in enumerate(items) if k == "pct"][0]
for c in range(2, 11):
    col = get_column_letter(c)
    cell = ws.cell(row=pct_row, column=c,
                   value=f'=COUNTIF({col}{chk_rows[0]}:{col}{chk_rows[-1]},"✔")/{len(chk_rows)}')
    cell.number_format = "0%"
    style(cell, font=F_BOLD, align=CENTER)
add_list_validation(ws, '"✔"', f"B{chk_rows[0]}:J{chk_rows[-1]}")
ws.freeze_panes = "B6"

# =====================================================================
# 7. 投稿後振り返り
# =====================================================================
ws = wb.create_sheet("07_投稿後振り返り")
setup(ws, "投稿後の振り返り（1本20〜30分・絶対に飛ばさない）",
      "詳しく書くほどパターンが見える。「もっと時間があったら」欄は渾身の1本のヒントになる。数値は公開7日後など同じタイミングで記録。",
      [5, 30, 16, 30, 30, 30, 26, 30, 32, 12, 12, 10, 10, 12, 30])
cols = ["#", "動画タイトル", "使ったレシピ", "学習目標は達成できた？何がわかった？", "うまくいったこと",
        "次の動画で重点的に改善すること", "新しく習得したスキル", "視聴者の反応（コメント等）",
        "もっと時間があったら何をしていた？", "再生数\n(7日)", "CTR(%)", "平均視聴率(%)", "登録者増",
        "感覚評価\n(1-5)", "気づき・パターン"]
header(ws, 6, cols)
example = ["IELTSライティング 独学で6.5→7.0…", "IELTS独学ラボ", "冒頭で結果を見せると30秒維持率+8%",
           "画面録画は撮影が早い", "中盤のテンポ（カットを増やす）", "字幕テンプレ作成",
           "「添削例をもっと見たい」多数", "実際の添削ビフォーアフターを5本分見せたかった",
           1200, 0.052, 0.41, 35, 4, "解説系は冒頭の結果提示が効く"]
input_block(ws, 7, 10, range(2, 16), example)
for r in range(7, 17):
    style(ws.cell(row=r, column=1, value=r - 6 if r < 16 else "渾身"), align=CENTER)
    ws.cell(row=r, column=10).number_format = "#,##0"
    ws.cell(row=r, column=11).number_format = "0.0%"
    ws.cell(row=r, column=12).number_format = "0.0%"
    ws.row_dimensions[r].height = 48
add_list_validation(ws, "'03_動画レシピ'!$C$6:$E$6", "C7:C16")
add_list_validation(ws, '"1,2,3,4,5"', "N7:N16")
ws.cell(row=5, column=10, value="CTR・視聴率は小数で（5.2%→0.052）").font = F_SUB

section(ws, 18, "レシピ別の平均（スプリント8本・自動）", 15)
header(ws, 19, ["", "レシピ", "本数", "平均再生数", "平均CTR", "平均視聴率", "平均登録者増", "平均感覚評価"])
for i in range(3):
    r = 20 + i
    col = "CDE"[i]
    style(ws.cell(row=r, column=2, value=f"=IF('03_動画レシピ'!{col}6=\"\",\"\",'03_動画レシピ'!{col}6)"), font=F_BOLD)
    style(ws.cell(row=r, column=3, value=f'=IF(B{r}="","",COUNTIF($C$7:$C$15,B{r}))'), align=CENTER)
    for c, src in zip(range(4, 9), "JKLMN"):
        f = f'=IF(B{r}="","",IFERROR(AVERAGEIF($C$7:$C$15,B{r},${src}$7:${src}$15),"-"))'
        cell = ws.cell(row=r, column=c, value=f)
        style(cell, align=CENTER)
        cell.number_format = {"J": "#,##0", "K": "0.0%", "L": "0.0%", "M": "0.0", "N": "0.0"}[src]
ws.freeze_panes = "C7"

# =====================================================================
# 8. 渾身の1本
# =====================================================================
ws = wb.create_sheet("08_渾身の1本")
setup(ws, "W10〜W11 フェーズ3：渾身の特大ヒット作",
      "これまでの全行動はこの1本のため。スプリントの学びと「もっと時間があったら」を全投入し、質を極限まで追求する。",
      [40, 90])
header(ws, 5, ["考えること", "あなたの答え"])
hero = [
    ("【振り返りから集める】", None),
    ("スプリントで最も反応が良かったレシピ（自動参照→07下の表）", ""),
    ("振り返りの「もっと時間があったら」で繰り返し出た要素", ""),
    ("振り返りの「次に改善すること」で繰り返し出た要素", ""),
    ("視聴者の反応で多かった要望", ""),
    ("W4で「渾身候補」に回した大きなアイデア", ""),
    ("【企画】", None),
    ("渾身の1本のテーマ", ""),
    ("なぜこれがチャンネルを爆発させ得るか（根拠：実績コンテンツ・データ）", ""),
    ("ターゲット／価値／雰囲気／差別化（最終レシピ）", ""),
    ("タイトル案（10案以上出して選ぶ）", ""),
    ("サムネイル案（複数の構図・スタイル）", ""),
    ("冒頭30秒のフック", ""),
    ("ストーリー構成（好奇心を引っ張る要素）", ""),
    ("スプリント時より時間・予算をかけるポイント", ""),
    ("【スケジュール】", None),
    ("企画・台本 締切", ""),
    ("撮影 締切", ""),
    ("編集 締切", ""),
    ("公開日", ""),
    ("公開時の告知（どこで・何を）", ""),
]
for i, (q, _) in enumerate(hero):
    r = 6 + i
    if q.startswith("【"):
        section(ws, r, q, 2)
        continue
    style(ws.cell(row=r, column=1, value=q), font=F_BOLD)
    style(ws.cell(row=r, column=2), fill=FILL_INPUT)
    if "締切" in q or q == "公開日":
        ws.cell(row=r, column=2).number_format = "yyyy/mm/dd"
    ws.row_dimensions[r].height = 36
ws["B7"] = ('=IFERROR(INDEX(\'07_投稿後振り返り\'!B20:B22,MATCH(MAX(\'07_投稿後振り返り\'!D20:D22),'
            '\'07_投稿後振り返り\'!D20:D22,0))&"（平均再生数ベース）","振り返りを記入すると表示")')
style(ws["B7"])
ws.cell(row=28, column=1, value="※制作チェックは 06_制作シート の「渾身の1本」列、振り返りは 07 の「渾身」行を使う。").font = F_SUB

# =====================================================================
# 9. 最終振り返り
# =====================================================================
ws = wb.create_sheet("09_最終振り返り")
setup(ws, "W12 最終振り返り：勢いを失わずに次の章へ",
      "全振り返りを見直し、公開した動画を新鮮な目で見返す。量（スプリント）と質（渾身）の両極端を経験した今、今後のバランスを決める。",
      [44, 86])
header(ws, 5, ["考えること", "あなたの答え"])
final = [
    ("【数字】", None),
    ("チャレンジ中の公開本数（自動）", "='05_スプリント管理'!C4&\"本（スプリント）＋渾身1本\""),
    ("開始時→終了時の登録者数", ""),
    ("最も伸びた動画とその理由", ""),
    ("【量 vs 質】", None),
    ("スプリント（量を極限まで）でどう感じた？何ができた？", ""),
    ("渾身の1本（質を極限まで）でどう感じた？何ができた？", ""),
    ("今後のバランスはどうしたい？", ""),
    ("【今後の方針】", None),
    ("今後の投稿スケジュール（元に戻す／変更する）", ""),
    ("確定した動画レシピ（黄金律）", ""),
    ("自分のクリエイティブな表現で明確になったこと", ""),
    ("視聴者とのつながりで変わったこと", ""),
    ("次の90日でやること", ""),
]
for i, (q, v) in enumerate(final):
    r = 6 + i
    if q.startswith("【"):
        section(ws, r, q, 2)
        continue
    style(ws.cell(row=r, column=1, value=q), font=F_BOLD)
    c = ws.cell(row=r, column=2, value=v or None)
    style(c, fill=None if v else FILL_INPUT)
    ws.row_dimensions[r].height = 40

# =====================================================================
# 10. 考えることリスト
# =====================================================================
ws = wb.create_sheet("10_考えることリスト")
setup(ws, "あなたが考えないといけないこと 一覧",
      "動画中の問いをすべて抜き出したチェックリスト。答えを書く場所は右の「記入先シート」。",
      [8, 16, 70, 24, 10])
header(ws, 5, ["週", "カテゴリ", "問い", "記入先シート", "考えた✔"])
qlist = [
    ("W1", "決意", "90日後の目標は何か？（書き出す）", "01_決意表明"),
    ("W1", "決意", "誰に・どこで宣言するか？", "01_決意表明"),
    ("W2", "研究", "なぜこれが面白いと思ったのか？", "02_スワイプファイル"),
    ("W2", "研究", "なぜクリックしたのか？", "02_スワイプファイル"),
    ("W2", "研究", "なぜ開始30秒経っても見続けているのか？", "02_スワイプファイル"),
    ("W2", "研究", "音楽や照明によってどんな感情が湧くか？", "02_スワイプファイル"),
    ("W2", "研究", "どの時点で退屈・離脱したか、きっかけは何か？", "02_スワイプファイル"),
    ("W3", "レシピ", "誰のために動画を作っている？何に興味・悩みがある？普段何を見ている？", "03_動画レシピ"),
    ("W3", "レシピ", "視聴者は何を得られる？（行動・理解・孤独の緩和・娯楽・インスピレーション）", "03_動画レシピ"),
    ("W3", "レシピ", "その価値をどのフォーマットで届ける？", "03_動画レシピ"),
    ("W3", "レシピ", "動画はどんな気分にさせる？", "03_動画レシピ"),
    ("W3", "レシピ", "なぜ他の誰でもなく自分の動画を見るべき？", "03_動画レシピ"),
    ("W3", "レシピ", "自分が生まれつき得意なことは？スワイプで繰り返し出る要素は？", "03_動画レシピ"),
    ("W4", "アイデア", "このアイデアはレシピを洗練させるのに役立つか？", "04_アイデア8本"),
    ("W4", "アイデア", "大ヒットのきっかけになり得るか？（実績コンテンツがあるか）", "04_アイデア8本"),
    ("W4", "アイデア", "普段の2倍のペースで台本・撮影・編集・公開できるか？無理なら簡略化できるか？", "04_アイデア8本"),
    ("W5", "スプリント", "2回目の公開宣言をどこで何と言うか？", "01_決意表明"),
    ("W5-9", "制作", "この動画から学ぼうとしている、たった1つのことは何か？", "05_スプリント管理"),
    ("W5-9", "振り返り", "どのレシピを使った？何がうまくいった？", "07_投稿後振り返り"),
    ("W5-9", "振り返り", "次の動画で何を重点的に改善する？", "07_投稿後振り返り"),
    ("W5-9", "振り返り", "どんな新しいスキルを習得した？", "07_投稿後振り返り"),
    ("W5-9", "振り返り", "視聴者の反応はどうだった？", "07_投稿後振り返り"),
    ("W5-9", "振り返り", "もっと時間があったら何をしていた？", "07_投稿後振り返り"),
    ("W5-9", "共有", "進捗をどこで・どの頻度で共有する？", "05_スプリント管理"),
    ("W10-11", "渾身", "スプリントの学びをどう1本に全投入するか？", "08_渾身の1本"),
    ("W12", "総括", "量と質、それぞれ重視した時にどう感じたか？", "09_最終振り返り"),
    ("W12", "総括", "今後はどんなバランスでいきたいか？", "09_最終振り返り"),
    ("W12", "総括", "投稿スケジュールを元に戻す？変更する？", "09_最終振り返り"),
]
for i, row in enumerate(qlist):
    r = 6 + i
    for c, v in enumerate(row, 1):
        style(ws.cell(row=r, column=c, value=v), align=CENTER if c in (1, 2) else WRAP)
    style(ws.cell(row=r, column=5), fill=FILL_INPUT, align=CENTER)
add_list_validation(ws, '"✔"', f"E6:E{5 + len(qlist)}")
ws.cell(row=4, column=3, value=f'=COUNTIF(E6:E{5 + len(qlist)},"✔")&" / {len(qlist)} 考えた"').font = F_BOLD
ws.freeze_panes = "A6"

out = "youtube_90day/YouTube90日チャレンジ_ワークブック.xlsx"
wb.save(out)
print("saved", out)

# =====================================================================
# 記入例入り版：全シートに例を入れて別ファイルで保存
# =====================================================================
import datetime as dt

from examples import EXAMPLES


def put(ws, ref, value):
    cell = ws[ref]
    cell.value = value
    cell.font = F_EX
    cell.fill = FILL_EX
    if isinstance(value, dt.date):
        cell.number_format = "yyyy/mm/dd"


for sheet, cells in EXAMPLES.items():
    for ref, value in cells.items():
        put(wb[sheet], ref, value)
for name in wb.sheetnames:
    wb[name]["A3"] = "記入例入り版：灰色の斜体＝記入例です。自分の内容に上書きしてください（空欄版は別ファイル）"
out = "youtube_90day/YouTube90日チャレンジ_記入例入り.xlsx"
wb.save(out)
print("saved", out)

# =====================================================================
# Googleドライブ用：準備編・制作編の2ファイルに分割（アップロードサイズ対策）
# =====================================================================
from openpyxl import load_workbook

PARTS = {
    "準備編": ["00_ロードマップ", "01_決意表明", "02_スワイプファイル", "03_動画レシピ", "04_アイデア8本", "10_考えることリスト"],
    "制作編": ["03_動画レシピ", "05_スプリント管理", "06_制作シート", "07_投稿後振り返り", "08_渾身の1本", "09_最終振り返り"],
}
for part, keep in PARTS.items():
    wbp = load_workbook(out)
    for name in wbp.sheetnames:
        if name not in keep:
            wbp.remove(wbp[name])
    if part == "準備編":
        ws = wbp["00_ロードマップ"]
        for r in (30, 31):
            ws.cell(row=r, column=3, value="→ 制作編ファイルで確認")
    else:
        wbp["03_動画レシピ"]["A2"] = "準備編で作ったレシピをここにも書き写してください（05〜08のプルダウンに使います）"
    wbp.save(f"youtube_90day/YouTube90日チャレンジ_{part}.xlsx")
    print("saved", part)
