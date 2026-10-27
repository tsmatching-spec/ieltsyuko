"""LINE公式アカウントの日別データ（CSV）と、その日のストーリーズ内容を「LINE登録数」シートに記入する。"""
import csv
import math
import glob
from datetime import date

from openpyxl import load_workbook

BOOK = "sns_tracking/SNS記録シート.xlsx"
START = date(2026, 10, 1)

wb = load_workbook(BOOK)
ws = wb["LINE登録数"]


def row_of(d):
    return 2 + (d - START).days


for path in sorted(glob.glob("sns_tracking/data/line_*.csv")):
    with open(path, encoding="utf-8") as f:
        for rec in csv.DictReader(f):
            r = row_of(date.fromisoformat(rec["日付"]))
            for col, key in ((3, "友だち登録数"), (4, "ブロック数"), (6, "有効友だち数"), (7, "累計友だち追加数")):
                ws.cell(row=r, column=col, value=int(rec[key]))

# その日のストーリーズ内容を📝欄に（add_story_images.py の STORIES を流用）
src = open("scripts/add_story_images.py", encoding="utf-8").read()
ns = {}
exec(src[src.index("STORIES = ["):src.index("]\n", src.index("STORIES = [")) + 2], ns)
by_day = {}
for day, seq, memo in ns["STORIES"]:
    by_day.setdefault(day, []).append(memo)
for day, memos in by_day.items():
    r = row_of(date(2026, 10, day))
    ws.cell(row=r, column=8, value="\n".join(f"{i}. {m}" for i, m in enumerate(memos, 1)))
    ws.cell(row=r, column=9, value=f"ストーリーズ{len(memos)}本")
    lines = sum(math.ceil((len(m) + 3) / 36) for m in memos)  # 1行に全角36字ほど
    ws.row_dimensions[r].height = 15 * lines + 4

wb.save(BOOK)
print("filled LINE data and post notes")
