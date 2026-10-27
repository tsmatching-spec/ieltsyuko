"""LINE登録数・投稿メモ / 月別ストーリーズ閲覧数 の記録用スプレッドシートを作成する。"""
import calendar
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

OUT = "sns_tracking/SNS記録シート.xlsx"
START = date(2026, 10, 1)
MONTHS = [(2026, 10), (2026, 11), (2026, 12), (2027, 1), (2027, 2), (2027, 3)]
STORY_SLOTS = 10  # 1日あたりのストーリーズ枠（横に並べる数）
STORY_FIRST_ROW = 5
ROWS_PER_DAY = 3  # スクショ / 内容 / 閲覧数
SLOT_COL = 6  # F列から横にスクショを並べる

HEAD_FILL = PatternFill("solid", fgColor="06C755")  # LINEグリーン
STORY_FILL = PatternFill("solid", fgColor="C13584")  # インスタ系
SUB_FILL = PatternFill("solid", fgColor="F3F3F3")
WEEKEND_FILL = PatternFill("solid", fgColor="FFF4E5")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
thin = Side(style="thin", color="CCCCCC")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center")
WRAP = Alignment(vertical="top", wrap_text=True)
WD = "月火水木金土日"


def header(ws, row, labels, fill):
    for col, label in enumerate(labels, 1):
        c = ws.cell(row=row, column=col, value=label)
        c.fill, c.font, c.alignment, c.border = fill, WHITE_BOLD, CENTER, BORDER


wb = Workbook()

# --- LINE登録数（毎日） ---
ws = wb.active
ws.title = "LINE登録数"
header(ws, 1, ["日付", "曜日", "LINE登録数", "累計", "📝 その日の投稿内容", "投稿先", "メモ"], HEAD_FILL)
end = date(*MONTHS[-1], calendar.monthrange(*MONTHS[-1])[1])
d, r = START, 2
while d <= end:
    ws.cell(row=r, column=1, value=d).number_format = "yyyy/mm/dd"
    ws.cell(row=r, column=2, value=WD[d.weekday()])
    ws.cell(row=r, column=4, value=f"=IF(C{r}=\"\",\"\",SUM($C$2:C{r}))")
    for col in range(1, 8):
        c = ws.cell(row=r, column=col)
        c.border = BORDER
        c.alignment = WRAP if col in (5, 7) else CENTER
        if d.weekday() >= 5:
            c.fill = WEEKEND_FILL
    d += timedelta(days=1)
    r += 1
last = r - 1
for col, w in zip("ABCDEFG", [12, 6, 12, 10, 50, 14, 30]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:G{last}"

# --- 月別まとめ ---
sm = wb.create_sheet("月別まとめ")
header(sm, 1, ["月", "LINE登録数（合計）", "1日平均", "ストーリーズ本数", "閲覧数（合計）", "閲覧数（平均）"], HEAD_FILL)
for i, (y, m) in enumerate(MONTHS, 2):
    first, lastday = date(y, m, 1), date(y, m, calendar.monthrange(y, m)[1])
    s = f"{m}月ストーリーズ"
    srange = f"'{s}'!{get_column_letter(SLOT_COL)}{STORY_FIRST_ROW}:{get_column_letter(SLOT_COL + STORY_SLOTS - 1)}{STORY_FIRST_ROW + lastday.day * ROWS_PER_DAY - 1}"
    slast = STORY_FIRST_ROW + lastday.day * ROWS_PER_DAY - 1
    sm.cell(row=i, column=1, value=f"{y}年{m}月")
    sm.cell(row=i, column=2, value=f"=SUMIFS('LINE登録数'!C:C,'LINE登録数'!A:A,\">=\"&DATE({y},{m},1),'LINE登録数'!A:A,\"<=\"&DATE({y},{m},{lastday.day}))")
    sm.cell(row=i, column=3, value=f"=ROUND(B{i}/{lastday.day},1)")
    sm.cell(row=i, column=4, value=f"=SUM('{s}'!C{STORY_FIRST_ROW}:C{slast})")
    sm.cell(row=i, column=5, value=f"=SUM('{s}'!D{STORY_FIRST_ROW}:D{slast})")
    sm.cell(row=i, column=6, value=f"=IFERROR(ROUND(E{i}/COUNT({srange}),0),\"\")")
    for col in range(1, 7):
        sm.cell(row=i, column=col).border = BORDER
        sm.cell(row=i, column=col).alignment = CENTER
for col, w in zip("ABCDEF", [12, 18, 10, 16, 16, 16]):
    sm.column_dimensions[col].width = w

# --- 月別ストーリーズ（1日=3行：スクショ / 内容 / 閲覧数、その日の投稿を横に並べる） ---
for y, m in MONTHS:
    st = wb.create_sheet(f"{m}月ストーリーズ")
    days = calendar.monthrange(y, m)[1]
    first, last = STORY_FIRST_ROW, STORY_FIRST_ROW + days * ROWS_PER_DAY - 1
    c0, c1 = get_column_letter(SLOT_COL), get_column_letter(SLOT_COL + STORY_SLOTS - 1)
    views = f"{c0}{first}:{c1}{last}"  # 数値は閲覧数の行にしか入らない
    st["A1"] = f"{y}年{m}月 ストーリーズ記録"
    st["A1"].font = Font(bold=True, size=14)
    for col, (label, formula) in enumerate([
        ("本数", f"=SUM(C{first}:C{last})"),
        ("閲覧数合計", f"=SUM(D{first}:D{last})"),
        ("閲覧数平均", f"=IFERROR(ROUND(SUM(D{first}:D{last})/COUNT({views}),0),\"\")"),
        ("最高閲覧数", f"=IF(COUNT({views})=0,\"\",MAX({views}))"),
    ], SLOT_COL):
        lc = st.cell(row=2, column=col, value=label)
        vc = st.cell(row=3, column=col, value=formula)
        lc.fill, lc.alignment, lc.border = SUB_FILL, CENTER, BORDER
        vc.font, vc.alignment, vc.border = Font(bold=True, size=13), CENTER, BORDER
    header(st, 4, ["日付", "曜日", "本数", "閲覧数合計", ""] + [f"{n}枚目" for n in range(1, STORY_SLOTS + 1)], STORY_FILL)
    for day in range(1, days + 1):
        d = date(y, m, day)
        img_r, memo_r, view_r = (first + (day - 1) * ROWS_PER_DAY + k for k in range(3))
        st.row_dimensions[img_r].height = 210
        st.row_dimensions[memo_r].height = 48
        st.row_dimensions[view_r].height = 20
        st.cell(row=img_r, column=1, value=d).number_format = "m/d"
        st.cell(row=img_r, column=2, value=WD[d.weekday()])
        st.cell(row=img_r, column=3, value=f"=COUNTA({c0}{memo_r}:{c1}{memo_r})")
        st.cell(row=img_r, column=4, value=f"=SUM({c0}{view_r}:{c1}{view_r})")
        for col in range(1, 5):
            st.merge_cells(start_row=img_r, start_column=col, end_row=view_r, end_column=col)
        for r, label in ((img_r, "📷"), (memo_r, "内容"), (view_r, "閲覧数")):
            st.cell(row=r, column=5, value=label)
            for col in range(1, SLOT_COL + STORY_SLOTS):
                c = st.cell(row=r, column=col)
                c.border = BORDER
                c.alignment = WRAP if r == memo_r and col >= SLOT_COL else CENTER
                if d.weekday() >= 5 and col < SLOT_COL:
                    c.fill = WEEKEND_FILL
            if r == memo_r:
                for col in range(SLOT_COL, SLOT_COL + STORY_SLOTS):
                    st.cell(row=r, column=col).font = Font(size=9)
            if r == view_r:
                for col in range(SLOT_COL, SLOT_COL + STORY_SLOTS):
                    st.cell(row=r, column=col).fill = PatternFill("solid", fgColor="FDF2F8")
        # 1日ごとの区切り線
        for col in range(1, SLOT_COL + STORY_SLOTS):
            c = st.cell(row=view_r, column=col)
            c.border = Border(left=thin, right=thin, top=thin, bottom=Side(style="medium", color="999999"))
    for col, w in zip("ABCDE", [7, 5, 6, 10, 7]):
        st.column_dimensions[col].width = w
    for col in range(SLOT_COL, SLOT_COL + STORY_SLOTS):
        st.column_dimensions[get_column_letter(col)].width = 23
    st.conditional_formatting.add(
        views,
        FormulaRule(formula=[f"AND(ISNUMBER({c0}{first}),{c0}{first}=MAX(${c0}${first}:${c1}${last}))"],
                    fill=PatternFill("solid", fgColor="FFE08A")),
    )
    st.freeze_panes = f"{c0}{first}"

wb.save(OUT)
print("saved", OUT)
