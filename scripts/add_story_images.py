"""切り出したストーリーズ画像を月別シートに順番に貼り付け、投稿日・内容メモを記入する。"""
from datetime import date

from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU

BOOK = "sns_tracking/SNS記録シート.xlsx"
SHEET = "10月ストーリーズ"
IMG_DIR = "sns_tracking/stories/2026-10"
FIRST_ROW = 5
IMG_HEIGHT = 260  # px（行の高さ220pt≒293pxに余白付きで収まるサイズ）
PAD = 8  # px（セルの左上からの余白）

# (日, その日の何枚目か, 内容メモ) を投稿順に。画像は IMG_DIR/10-DD_NN.png
STORIES = [
    (1, 1, "ポッドキャスト更新 #3 負け癖からの脱却"),
    (1, 2, "YouTubeに届いたコメント（社会人は時間がないから伸びない？→答えはNO）"),
    (1, 3, "個別学習相談に来られる方々からの声（つまり…）"),
    (1, 4, "正しい方向に最短で進める仕組みが必要／IELTS個別学習相談リンク"),
]

wb = load_workbook(BOOK)
ws = wb[SHEET]
ws.column_dimensions["C"].width = 24
for n, (day, seq, memo) in enumerate(STORIES):
    row = FIRST_ROW + n
    ws.cell(row=row, column=2, value=date(2026, 10, day))
    ws.cell(row=row, column=4, value=memo)
    img = Image(f"{IMG_DIR}/10-{day:02d}_{seq:02d}.png")
    img.width, img.height = round(img.width * IMG_HEIGHT / img.height), IMG_HEIGHT
    marker = AnchorMarker(col=2, colOff=pixels_to_EMU(PAD), row=row - 1, rowOff=pixels_to_EMU(PAD))
    img.anchor = OneCellAnchor(_from=marker, ext=XDRPositiveSize2D(pixels_to_EMU(img.width), pixels_to_EMU(img.height)))
    ws.add_image(img)
wb.save(BOOK)
print(f"added {len(STORIES)} stories to {SHEET}")
