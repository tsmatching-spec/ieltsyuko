"""個別学習相談フォームの回答（Googleフォームのxlsx書き出し）から、指定日以降の申込者を記入する。

個人情報を含むため、入力ファイル・出力ファイルはgitに入れない（.gitignore 参照）。
使い方: python3 scripts/fill_consultations.py <フォーム回答.xlsx> [開始日 YYYY-MM-DD]
審査結果・成約は sns_tracking/private/相談結果.csv（お名前,審査結果,成約）から読む。
"""
import csv
import os
import sys
from collections import defaultdict
from datetime import date, datetime

from openpyxl import load_workbook

BOOK = "sns_tracking/SNS記録シート.xlsx"
RESULTS = "sns_tracking/private/相談結果.csv"
START = date(2026, 10, 1)

src = sys.argv[1]
since = date.fromisoformat(sys.argv[2]) if len(sys.argv) > 2 else START

form = load_workbook(src, data_only=True).active
apps = []
for ts, name, line_name, job, *_ in form.iter_rows(min_row=2, values_only=True):
    if not isinstance(ts, datetime) or ts.date() < since:
        continue
    clean = lambda v: " ".join(str(v or "").replace("　", " ").split())
    apps.append((ts, clean(name), clean(line_name), clean(job)))
apps.sort()

results = {}
if os.path.exists(RESULTS):
    with open(RESULTS, encoding="utf-8") as f:
        results = {row["お名前"]: row for row in csv.DictReader(f)}

wb = load_workbook(BOOK)
cs = wb["個別相談申込"]
for r, (ts, name, line_name, job) in enumerate(apps, 2):
    cs.cell(row=r, column=1, value=datetime(ts.year, ts.month, ts.day))
    cs.cell(row=r, column=2, value=ts.time().replace(microsecond=0))
    cs.cell(row=r, column=3, value=name)
    cs.cell(row=r, column=4, value=line_name)
    res = results.get(name, {})
    cs.cell(row=r, column=5, value=res.get("審査結果") or "審査中")
    cs.cell(row=r, column=6, value=res.get("成約") or None)
    cs.cell(row=r, column=7, value=job)

ws = wb["LINE登録数"]
by_day = defaultdict(list)
for ts, name, *_ in apps:
    by_day[ts.date()].append(name)
for d, names in by_day.items():
    ws.cell(row=2 + (d - START).days, column=12, value="\n".join(names))

wb.save(BOOK)
print(f"added {len(apps)} applications since {since}")
